from copy import copy

import openpyxl
import pytest

from ec101.contract.schema import SHEET_COLUMNS, STANDARD_SHEETS
from ec101.contract.workbook import read_standard_workbook, write_standard_workbook


def manifest_row():
    return {
        "模板版本": "1.0",
        "转换工具版本": "test",
        "经销商名称": "测试经销商",
        "平台名称": "测试平台",
        "数据开始日期": "2026-09-01",
        "数据结束日期": "2026-09-30",
        "生成时间": "2026-10-08 12:00:00",
    }


def empty_tables():
    return {sheet: [] for sheet in STANDARD_SHEETS if sheet != "导入清单"}


def test_round_trip_skips_blank_rows_and_preserves_long_identifiers(tmp_path):
    path = tmp_path / "standard.xlsx"
    tables = empty_tables()
    order_id = "12345678901234567890"
    coupon_id = "98765432109876543210"
    tables["标准订单明细"] = [{"单据编号": order_id}]
    tables["标准优惠券核销明细"] = [{"优惠券编号": coupon_id}]
    write_standard_workbook(path, tables, manifest_row())

    workbook = openpyxl.load_workbook(path)
    workbook["标准订单明细"].append([None] * len(SHEET_COLUMNS["标准订单明细"]))
    workbook.save(path)
    workbook.close()

    restored = read_standard_workbook(path)

    assert restored.manifest == manifest_row()
    assert restored.tables["标准订单明细"] == [{
        column: order_id if column == "单据编号" else ""
        for column in SHEET_COLUMNS["标准订单明细"]
    }]
    assert restored.tables["标准优惠券核销明细"][0]["优惠券编号"] == coupon_id


def test_read_rejects_reordered_sheets(tmp_path):
    path = tmp_path / "reordered.xlsx"
    write_standard_workbook(path, empty_tables(), manifest_row())
    workbook = openpyxl.load_workbook(path)
    workbook._sheets[0], workbook._sheets[1] = workbook._sheets[1], workbook._sheets[0]
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="工作表顺序"):
        read_standard_workbook(path)


def test_read_rejects_mismatched_headers(tmp_path):
    path = tmp_path / "bad-header.xlsx"
    write_standard_workbook(path, empty_tables(), manifest_row())
    workbook = openpyxl.load_workbook(path)
    workbook["标准客户"].cell(row=1, column=1, value="错误表头")
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="表头"):
        read_standard_workbook(path)


@pytest.mark.parametrize("manifest", [[], [manifest_row(), copy(manifest_row())]])
def test_manifest_must_have_exactly_one_row(tmp_path, manifest):
    path = tmp_path / "bad-manifest.xlsx"

    with pytest.raises(ValueError, match="导入清单.*一行"):
        write_standard_workbook(path, empty_tables(), manifest)
