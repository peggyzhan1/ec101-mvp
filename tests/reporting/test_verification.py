from openpyxl import load_workbook

from ec101.reporting.verification import build_verification_workbook


def test_verification_workbook_contains_result_sheets_and_reconciles_unlinked_coupon(reporting_db, tmp_path):
    db, ids = reporting_db
    output = build_verification_workbook(db, ids["run_id"], tmp_path / "verification.xlsx")

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["核算信息", "活动费用汇总", "优惠券费用汇总", "释放候选", "未关联券"]
    assert workbook["核算信息"]["B1"].value == ids["run_id"]
    coupon_sheet = workbook["优惠券费用汇总"]
    assert coupon_sheet["A2"].value == 1
    assert coupon_sheet["G2"].value == 1
    assert coupon_sheet["H2"].value == 3
    unlinked_sheet = workbook["未关联券"]
    assert unlinked_sheet.max_row == 2
    assert unlinked_sheet["E2"].value == 3
    workbook.close()
