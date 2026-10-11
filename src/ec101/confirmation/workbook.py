from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import math
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


RECIPIENT_HEADERS = (
    "费用周期",
    "太古客户编号",
    "太古客户名称",
    "应发金额",
    "备注",
)

SURVEY_HEADERS = (
    "费用周期",
    "太古客户编号",
    "太古客户名称",
    "是否收到本期全部协议费用",
    "未收到原因",
    "补充说明",
    "答卷手机号",
    "提交时间",
    "问卷星原始序号",
)


class FeeRecipientValidationError(ValueError):
    pass


class SurveyValidationError(ValueError):
    pass


def normalize_customer_code(value: Any) -> str:
    if value is None or isinstance(value, bool):
        raise FeeRecipientValidationError("太古客户编号不能为空")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer() or abs(value) >= 10**15:
            raise FeeRecipientValidationError("太古客户编号无法无损转换，请在 Excel 中设为文本")
        return str(int(value))
    text = str(value).strip()
    if not text:
        raise FeeRecipientValidationError("太古客户编号不能为空")
    if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", text):
        raise FeeRecipientValidationError("太古客户编号不能使用科学计数法，请在 Excel 中设为文本")
    return text


def parse_positive_amount(value: Any) -> Decimal:
    if value is None or isinstance(value, bool) or str(value).strip() == "":
        raise FeeRecipientValidationError("应发金额不能为空")
    try:
        amount = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise FeeRecipientValidationError("应发金额必须是数字") from exc
    if not amount.is_finite() or amount <= 0:
        raise FeeRecipientValidationError("应发金额必须大于零")
    if amount.as_tuple().exponent < -2:
        raise FeeRecipientValidationError("应发金额最多保留两位小数")
    return amount.quantize(Decimal("0.01"))


def read_fee_recipient_workbook(path: str | Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise FeeRecipientValidationError("应发名单必须且只能包含一个工作表")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple("" if value is None else str(value).strip() for value in next(rows, ()))
        if headers != RECIPIENT_HEADERS:
            raise FeeRecipientValidationError("应发名单表头必须与标准字段完全一致")
        result = []
        for row_number, values in enumerate(rows, start=2):
            padded = tuple(values) + (None,) * (len(RECIPIENT_HEADERS) - len(values))
            values = padded[:len(RECIPIENT_HEADERS)]
            if not any(value is not None and str(value).strip() for value in values):
                continue
            try:
                code = normalize_customer_code(values[1])
                amount = parse_positive_amount(values[3])
            except FeeRecipientValidationError as exc:
                raise FeeRecipientValidationError(f"第 {row_number} 行：{exc}") from exc
            result.append({
                "费用周期": "" if values[0] is None else str(values[0]).strip(),
                "太古客户编号": code,
                "太古客户名称": "" if values[2] is None else str(values[2]).strip(),
                "应发金额": format(amount, ".2f"),
                "备注": "" if values[4] is None else str(values[4]).strip(),
            })
        return result
    finally:
        workbook.close()


def _survey_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def read_survey_workbook(path: str | Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise SurveyValidationError("问卷答卷必须且只能包含一个工作表")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(_survey_text(value) for value in next(rows, ()))
        if headers != SURVEY_HEADERS:
            raise SurveyValidationError("问卷答卷表头必须与标准字段完全一致")
        result = []
        for row_number, values in enumerate(rows, start=2):
            padded = tuple(values) + (None,) * (len(SURVEY_HEADERS) - len(values))
            values = padded[:len(SURVEY_HEADERS)]
            if not any(_survey_text(value) for value in values):
                continue
            code = ""
            if _survey_text(values[1]):
                try:
                    code = normalize_customer_code(values[1])
                except FeeRecipientValidationError as exc:
                    raise SurveyValidationError(f"第 {row_number} 行：{exc}") from exc
            normalized = [_survey_text(value) for value in values]
            normalized[1] = code
            result.append(dict(zip(SURVEY_HEADERS, normalized)))
        return result
    finally:
        workbook.close()
