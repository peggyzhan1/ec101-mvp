from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from .schema import SHEET_COLUMNS, STANDARD_SHEETS


@dataclass(frozen=True)
class StandardWorkbook:
    tables: dict[str, list[dict[str, str]]]
    manifest: dict[str, str]


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _text_cells(worksheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(worksheet, value=_text(value))
        cell.data_type = "s"
        cell.number_format = "@"
        cells.append(cell)
    return cells


def write_standard_workbook(
    path: str | Path,
    tables: Mapping[str, Sequence[Mapping[str, Any]]],
    manifest: Mapping[str, Any],
) -> None:
    if not isinstance(manifest, Mapping):
        raise ValueError("导入清单必须且只能包含一行")
    unknown_sheets = set(tables) - set(STANDARD_SHEETS[1:])
    if unknown_sheets:
        raise ValueError(f"包含未知工作表：{', '.join(sorted(unknown_sheets))}")

    workbook = Workbook(write_only=True)
    for sheet in STANDARD_SHEETS:
        worksheet = workbook.create_sheet(sheet)
        columns = SHEET_COLUMNS[sheet]
        worksheet.append(_text_cells(worksheet, columns))
        rows = [manifest] if sheet == "导入清单" else tables.get(sheet, ())
        for row in rows:
            worksheet.append(_text_cells(worksheet, (row.get(column) for column in columns)))

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def read_standard_workbook(path: str | Path, *, validate_contract: bool = False) -> StandardWorkbook:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if validate_contract and tuple(workbook.sheetnames) != STANDARD_SHEETS:
            raise ValueError("标准工作簿的工作表顺序不正确")

        loaded: dict[str, list[dict[str, str]]] = {}
        for sheet in STANDARD_SHEETS:
            worksheet = workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            actual_headers = tuple(_text(value) for value in next(rows, ()))
            expected_headers = SHEET_COLUMNS[sheet]
            if validate_contract and actual_headers != expected_headers:
                raise ValueError(f"{sheet} 的表头与标准契约不一致")

            records: list[dict[str, str]] = []
            for values in rows:
                normalized = tuple(_text(value) for value in values)
                normalized += ("",) * (len(expected_headers) - len(normalized))
                normalized = normalized[:len(expected_headers)]
                if not any(value.strip() for value in normalized):
                    continue
                records.append(dict(zip(expected_headers, normalized)))
            loaded[sheet] = records

        manifest_rows = loaded.pop("导入清单")
        if validate_contract and len(manifest_rows) != 1:
            raise ValueError("导入清单必须且只能包含一行")
        return StandardWorkbook(tables=loaded, manifest=manifest_rows[0] if manifest_rows else {})
    finally:
        workbook.close()
