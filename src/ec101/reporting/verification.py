from pathlib import Path
import sqlite3

from openpyxl import Workbook


def _append_rows(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append([row[column] for column in headers])


def build_verification_workbook(db: sqlite3.Connection, calculation_run_id: int, output_path: Path) -> Path:
    run = db.execute(
        "SELECT calculation_run_id, import_batch_id, calc_date, release_cutoff, activity_benefit, "
        "coupon_benefit, total_benefit, released_activity_benefit, released_coupon_benefit, "
        "participating_orders, released_orders, status FROM calculation_run WHERE calculation_run_id = ?",
        (calculation_run_id,),
    ).fetchone()
    if run is None:
        raise LookupError(f"核算结果 {calculation_run_id} 不存在")
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "核算信息"
    metadata.append(["核算批次", run["calculation_run_id"]])
    for key in ("import_batch_id", "calc_date", "release_cutoff", "activity_benefit", "coupon_benefit", "total_benefit", "released_activity_benefit", "released_coupon_benefit", "participating_orders", "released_orders", "status"):
        metadata.append([key, run[key]])

    activities = [dict(row) for row in db.execute(
        "SELECT afs.activity_id AS activityId, a.activity_no AS activityNo, a.activity_name AS activityName, "
        "afs.actual_discount_total AS actualDiscountTotal, afs.released_orders AS releasedOrders, "
        "afs.released_amount AS releasedAmount, afs.pending_orders AS pendingOrders, afs.pending_amount AS pendingAmount "
        "FROM activity_fee_summary afs JOIN activity a ON a.activity_id = afs.activity_id "
        "WHERE afs.calculation_run_id = ? ORDER BY afs.activity_fee_summary_id", (calculation_run_id,)
    ).fetchall()]
    sheet = workbook.create_sheet("活动费用汇总")
    _append_rows(sheet, list(activities[0]) if activities else ["activityId"], activities)

    coupons = [dict(row) for row in db.execute(
        "SELECT coupon_config_id AS couponConfigId, used_count AS usedCount, used_amount AS usedAmount, "
        "linked_amount AS linkedAmount, releasable_count AS releasableCount, releasable_amount AS releasableAmount, "
        "unlinked_count AS unlinkedCount, unlinked_amount AS unlinkedAmount, settle_status AS settleStatus "
        "FROM coupon_fee_summary WHERE calculation_run_id = ? ORDER BY coupon_fee_summary_id", (calculation_run_id,)
    ).fetchall()]
    sheet = workbook.create_sheet("优惠券费用汇总")
    _append_rows(sheet, list(coupons[0]) if coupons else ["couponConfigId"], coupons)

    candidates = [dict(row) for row in db.execute(
        "SELECT oh.order_no AS orderNo, rc.is_candidate AS isCandidate, rc.reason, "
        "rc.activity_benefit AS activityBenefit, rc.coupon_benefit AS couponBenefit "
        "FROM release_candidate rc JOIN order_header oh ON oh.order_id = rc.order_id "
        "WHERE rc.calculation_run_id = ? ORDER BY rc.release_candidate_id", (calculation_run_id,)
    ).fetchall()]
    sheet = workbook.create_sheet("释放候选")
    _append_rows(sheet, list(candidates[0]) if candidates else ["orderNo"], candidates)

    unlinked = [dict(row) for row in db.execute(
        "SELECT cr.coupon_no AS couponNo, cr.customer_no AS customerNo, cr.customer_name AS customerName, "
        "cr.status, cr.discount_amount AS discountAmount FROM coupon_redemption cr "
        "JOIN calculation_run run ON run.import_batch_id = cr.import_batch_id "
        "WHERE run.calculation_run_id = ? AND cr.order_id IS NULL ORDER BY cr.coupon_redemption_id",
        (calculation_run_id,),
    ).fetchall()]
    sheet = workbook.create_sheet("未关联券")
    _append_rows(sheet, list(unlinked[0]) if unlinked else ["couponNo"], unlinked)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
